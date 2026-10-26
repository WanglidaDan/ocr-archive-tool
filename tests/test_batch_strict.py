import sys, unittest, tempfile, zipfile, re
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
sys.path.insert(0, str(Path(__file__).resolve().parent))
from batch_export import describe_batch, export_separate
from extraction import extract_details
from fixtures import example_template

TEMPLATE = example_template()


class StrictBatchTests(unittest.TestCase):
    def test_per_document_classification_and_missing(self):
        for value in [
            "内部",
            "普通商密",
            "非涉密",
            "无密级",
            "油商密★",
            "油商密★★",
            "石油内部",
            "石油商密",
            "秘密",
            "机密",
            "绝密",
            "",
        ]:
            r = extract_details(value + "\n文件内容", ["密级"], 1)
            self.assertEqual(r[0]["密级"], value)
        self.assertEqual(
            extract_details("正文提到内部管理", ["密级"], 1)[0]["密级"], ""
        )
        self.assertEqual(
            extract_details("无标记", ["密级"], 1, additional={"密级": "内部"})[0][
                "密级"
            ],
            "",
        )

    def test_classification_conflict_and_position(self):
        v, e, w = extract_details(
            "--- 第 1 页 ---\n内部\n--- 第 2 页 ---\n普通商密", ["密级"], 2
        )
        self.assertEqual(v["密级"], "")
        self.assertEqual(e["密级"]["status"], "候选冲突")
        rows = [
            {
                "text": "内部",
                "page": 1,
                "confidence": 1,
                "boxes": [[[10, 700], [50, 700], [50, 720], [10, 720]]],
                "page_width": 600,
                "page_height": 800,
            }
        ]
        self.assertEqual(extract_details("", ["密级"], 1, rows)[0]["密级"], "")

    def test_one_workbook_per_file_and_no_overwrite(self):
        from openpyxl import load_workbook

        records = [
            {
                "path": "甲/同名.pdf",
                "values": {"题名": "第一份", "密级": "内部", "正本数量": "1"},
            },
            {
                "path": "乙/同名.pdf",
                "values": {"题名": "第二份", "密级": "普通商密", "正本数量": "1"},
            },
        ]
        with tempfile.TemporaryDirectory() as tmp:
            results = export_separate(records, tmp, TEMPLATE, "管理类_件")
            self.assertEqual(len(results), 2)
            self.assertTrue(all(not r["error"] for r in results))
            self.assertEqual(Path(results[0]["output"]).name, "同名.xlsx")
            self.assertEqual(Path(results[1]["output"]).name, "同名_2.xlsx")
            self.assertNotEqual(results[0]["output"], results[1]["output"])
            for record, result in zip(records, results):
                b = load_workbook(result["output"])
                self.assertEqual(b["管理类_件"]["D2"].value, record["values"]["题名"])
                self.assertIsNone(b["管理类_件"]["D3"].value)
                from copy import copy

                original = load_workbook(TEMPLATE)
                for x, y in zip(original["管理类_件"][2], b["管理类_件"][2]):
                    for attr in [
                        "font",
                        "fill",
                        "border",
                        "alignment",
                        "protection",
                        "number_format",
                    ]:
                        self.assertEqual(copy(getattr(x, attr)), copy(getattr(y, attr)))
                with (
                    zipfile.ZipFile(TEMPLATE) as a,
                    zipfile.ZipFile(result["output"]) as z,
                ):
                    self.assertEqual(a.namelist(), z.namelist())
                    self.assertEqual(
                        [n for n in a.namelist() if a.read(n) != z.read(n)],
                        ["xl/worksheets/sheet1.xml"],
                    )
                    # Remove only the written data cells; every other byte of the
                    # changed worksheet, including headers and validators, matches.
                    aa = a.read("xl/worksheets/sheet1.xml").decode()
                    zz = z.read("xl/worksheets/sheet1.xml").decode()
                    for address in ["D2", "K2", "T2"]:
                        pattern = r'<c\b[^>]*\br="' + address + r'"[^>]*(?:/>|>.*?</c>)'
                        aa = re.sub(pattern, "", aa, flags=re.S)
                        zz = re.sub(pattern, "", zz, flags=re.S)
                    self.assertEqual(aa, zz)
            old = Path(results[0]["output"]).read_bytes()
            more = export_separate(records[:1], tmp, TEMPLATE, "管理类_件")
            self.assertEqual(Path(more[0]["output"]).name, "同名_3.xlsx")
            self.assertNotEqual(more[0]["output"], results[0]["output"])
            self.assertEqual(Path(results[0]["output"]).read_bytes(), old)

    def test_bad_file_does_not_stop_other_files(self):
        records = [
            {"path": "bad.pdf", "values": {"密级": "未知密级"}},
            {"path": "good.pdf", "values": {"密级": "非涉密"}},
        ]
        with tempfile.TemporaryDirectory() as tmp:
            results = export_separate(records, tmp, TEMPLATE, "管理类_件")
            self.assertTrue(results[0]["error"])
            self.assertFalse(results[1]["error"])
            self.assertEqual(Path(results[1]["output"]).name, "good.xlsx")
            self.assertEqual(
                [path.name for path in Path(tmp).glob("*.xlsx")], ["good.xlsx"]
            )
            summary = describe_batch(
                [
                    {**records[0], "status": "待核对", "text": "x"},
                    {**records[1], "status": "已核对", "text": "y"},
                    {"path": "miss.pdf", "status": "处理失败", "text": ""},
                ],
                results,
            )
            self.assertIn("成功 1 份，失败 2 份，待核对 0 份。", summary)
            self.assertIn("bad.pdf：", summary)
            self.assertIn("miss.pdf：尚未识别成功，未导出", summary)
            self.assertNotIn("待核对：", summary)

    def test_summary_counts_review_separately_from_failure(self):
        text = describe_batch(
            [
                {"path": "a.pdf", "status": "待核对", "text": "a"},
                {"path": "b.pdf", "status": "已核对", "text": "b"},
            ],
            [
                {"source": "a.pdf", "output": "a.xlsx", "error": ""},
                {"source": "b.pdf", "output": "b.xlsx", "error": ""},
            ],
        )
        self.assertIn("成功 2 份，失败 0 份，待核对 1 份。", text)
        self.assertIn("待核对：a.pdf", text)
        self.assertNotIn("b.pdf", text.split("待核对：", 1)[-1])

    def test_stop_before_start_writes_nothing(self):
        with tempfile.TemporaryDirectory() as tmp:
            self.assertEqual(
                export_separate(
                    [{"path": "a.pdf", "values": {}}],
                    tmp,
                    TEMPLATE,
                    "管理类_件",
                    cancelled=lambda: True,
                ),
                [],
            )
            self.assertEqual(list(Path(tmp).iterdir()), [])


if __name__ == "__main__":
    unittest.main()
