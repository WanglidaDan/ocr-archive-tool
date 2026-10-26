import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
sys.path.insert(0, str(Path(__file__).resolve().parent))
from core import extract, export_xlsx, DEFAULT_FIELDS, Processor, Cancelled, local_llm
from fixtures import example_scan


class Tests(unittest.TestCase):
    def test_rules_and_ambiguity(self):
        text = "文件名称：档案归档通知\n文号：档发〔2026〕12号\n日期：2026年10月26日\n责任单位：档案馆"
        out = extract(text, DEFAULT_FIELDS, 2)
        self.assertEqual(out["文号"], "档发〔2026〕12号")
        self.assertEqual(out["页数"], "2")
        self.assertEqual(out["责任单位"], "档案馆")
        self.assertEqual(
            extract("2026年10月26日\n2025年1月1日", ["日期"], 1)["日期"], ""
        )
        self.assertEqual(extract("没有任何字段", ["文号"], 1)["文号"], "")

    def test_excel_formula_and_raw(self):
        from openpyxl import load_workbook

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "test.xlsx"
            export_xlsx(
                path,
                [
                    {
                        "path": "test.png",
                        "status": "待核对",
                        "values": {"文件名称": "=1+1"},
                        "text": "x" * 65000,
                    }
                ],
                ["文件名称"],
            )
            book = load_workbook(path)
            self.assertEqual(book.active["A2"].data_type, "s")
            self.assertEqual(book.active["A2"].value, "=1+1")
            self.assertEqual(book["OCR原文"].max_row, 4)
            self.assertEqual(
                "".join(book["OCR原文"].cell(n, 3).value for n in range(2, 5)),
                "x" * 65000,
            )

    def test_model_evidence_filter(self):
        import json
        from unittest.mock import patch, MagicMock

        response = MagicMock()
        response.__enter__.return_value.read.return_value = json.dumps(
            {
                "message": {
                    "content": json.dumps(
                        {"责任单位": "档案馆", "日期": "2029年1月1日"}
                    )
                }
            }
        ).encode()
        opener = MagicMock()
        opener.open.return_value = response
        with patch("urllib.request.build_opener", return_value=opener):
            values = local_llm("责任单位：档案馆", ["责任单位", "日期"], "local-model")
        self.assertEqual(values, {"责任单位": "档案馆", "日期": ""})
        request = opener.open.call_args.args[0]
        self.assertEqual(request.full_url, "http://127.0.0.1:11434/api/chat")
        self.assertEqual(extract("文号：甲号\n文号：乙号", ["文号"], 2)["文号"], "")

    def test_cancel(self):
        with self.assertRaises(Cancelled):
            Processor().recognize(example_scan(), DEFAULT_FIELDS, lambda: True)


if __name__ == "__main__":
    unittest.main()
