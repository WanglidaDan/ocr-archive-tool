import re
import subprocess
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
sys.path.insert(0, str(Path(__file__).resolve().parent))
from openpyxl import load_workbook
from batch_export import export_mapped_separate, write_summary
from catalog_defaults import (
    DEFAULTS,
    DEFAULT_SHEET,
    effective_defaults,
    validate_defaults,
)
from catalog_map import format_file_size, map_piece
from fixtures import (
    INNER_HEADERS,
    PIECE_HEADERS,
    VOLUME_HEADERS,
    example_template,
    fake_units,
)
from template_io import inspect_template
from tools.compare_template import compare_workbooks
from tools.require_template import require_standard_template

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = example_template()


class FakeTemplateTests(unittest.TestCase):
    def test_shape_uses_fake_units_and_named_ranges(self):
        book = load_workbook(TEMPLATE)
        self.assertEqual(
            book.sheetnames[:3],
            ["管理类_件", "管理类_卷", "管理类_卷内件"],
        )
        for title in ("全宗号", "类别号", "保管期限", "密级", "模版类型"):
            self.assertIn(title, book.sheetnames)
        piece = book["管理类_件"]
        self.assertEqual([cell.value for cell in piece[1]], PIECE_HEADERS)
        self.assertEqual([cell.value for cell in book["管理类_卷"][1]], VOLUME_HEADERS)
        self.assertEqual(
            [cell.value for cell in book["管理类_卷内件"][1]], INNER_HEADERS
        )
        by_sqref = {}
        for dv in piece.data_validations.dataValidation:
            by_sqref.setdefault(str(dv.sqref), dv)
        self.assertEqual(by_sqref["D1"].formula1, '"题名"')
        self.assertEqual(by_sqref["G2:G100001"].formula1, "=FilingUnit")
        self.assertEqual(by_sqref["B2:B100001"].formula1, "=OUID")
        self.assertEqual(by_sqref["V2:V100001"].formula1, "1900")
        self.assertIn("100001", str(by_sqref["F2:F100001"].sqref))
        self.assertEqual(
            book.defined_names["FilingUnit"].attr_text, "'归档单位'!$A$1:$A$60"
        )
        self.assertEqual(
            book.defined_names["OUID"].attr_text, "'归档单位(系统生成)'!$A$1:$A$60"
        )
        self.assertEqual(
            book.defined_names["RetentionPeriod"].attr_text, "'保管期限'!$A$1:$A$3"
        )
        meta = inspect_template(TEMPLATE, "管理类_件")
        self.assertEqual(meta["choices"]["保管期限"], ["永久", "30年", "10年"])
        self.assertEqual(meta["choices"]["模版类型"], ["收文", "发文"])
        self.assertIn("普通商密", meta["choices"]["密级"])
        self.assertIn("非涉密", meta["choices"]["密级"])
        self.assertNotIn("内部", meta["choices"]["密级"])
        self.assertEqual(meta["choices"]["归档单位"][:2], ["示例单位A", "示例单位B"])
        self.assertEqual(meta["choices"]["归档单位"], fake_units())
        for options in meta["choices"].values():
            for option in options:
                for banned in ("档案馆", "办公室", "公司"):
                    self.assertNotIn(banned, option)
        book.close()

    def test_mapping_and_builtin_defaults(self):
        self.assertEqual(DEFAULT_SHEET, "管理类_件")
        self.assertEqual(DEFAULTS["保管期限"], "30年")
        self.assertEqual(DEFAULTS["载体类型"], "纸质")
        self.assertTrue(
            all(
                value == ""
                for field, value in DEFAULTS.items()
                if field not in ("保管期限", "载体类型")
            )
        )
        meta = inspect_template(TEMPLATE, "管理类_件")
        with tempfile.TemporaryDirectory() as tmp:
            scan = Path(tmp) / "扫描件.png"
            scan.write_bytes(b"12345")
            record = {
                "path": str(scan),
                "values": {
                    "题名": "关于归档的通知",
                    "责任者": "示例单位A",
                    "文件日期": "2026年10月26日",
                    "页数": "4",
                    "正本数量": "1",
                    "文号": "档发〔2026〕12号",
                    "密级": "普通商密",
                },
            }
            mapped = map_piece(record, meta["choices"], 3)
            self.assertEqual(mapped["题名"], "关于归档的通知")
            self.assertEqual(mapped["责任者"], "示例单位A")
            self.assertEqual(mapped["文件日期"], "2026-10-26")
            self.assertEqual(mapped["年度"], "2026")
            self.assertEqual(mapped["文件编号"], "档发〔2026〕12号")
            self.assertEqual(mapped["页数"], "4")
            self.assertEqual(mapped["正本数量"], "1")
            self.assertEqual(mapped["密级"], "普通商密")
            self.assertEqual(mapped["归档单位"], "示例单位A")
            self.assertEqual(mapped["归档单位(系统生成)"], "示例单位A")
            self.assertEqual(mapped["临时序号"], "3")
            self.assertEqual(mapped["电子文件数量"], "1")
            self.assertEqual(mapped["电子文件关联项"], "扫描件.png")
            self.assertEqual(mapped["电子文件大小"], format_file_size(5))
            self.assertEqual(mapped["电子文件大小"], "1")
            self.assertEqual(
                map_piece(record, meta["choices"], 3, size_unit="B")["电子文件大小"],
                "5",
            )
            larger = Path(tmp) / "较大.png"
            larger.write_bytes(b"x" * 2048)
            self.assertEqual(
                map_piece({"path": str(larger), "values": {}}, meta["choices"], 1)[
                    "电子文件大小"
                ],
                "1",
            )
            self.assertEqual(
                map_piece(
                    {"path": str(larger), "values": {}},
                    meta["choices"],
                    1,
                    size_unit="KB",
                )["电子文件大小"],
                "2",
            )
            self.assertEqual(mapped["保管期限"], "30年")
            self.assertEqual(mapped["载体类型"], "纸质")
            for field, value in DEFAULTS.items():
                if not value:
                    self.assertNotIn(field, mapped)
            unknown = map_piece(
                {
                    "path": str(scan),
                    "values": {"密级": "未定义密级", "责任者": "不存在的单位"},
                },
                meta["choices"],
                1,
            )
            self.assertNotIn("密级", unknown)
            self.assertNotIn("归档单位", unknown)
            self.assertNotIn("归档单位(系统生成)", unknown)
            manual = map_piece(
                {
                    "path": str(scan),
                    "values": {"责任者": "不存在的单位", "归档单位": "示例单位B"},
                },
                meta["choices"],
                1,
            )
            self.assertEqual(manual["归档单位"], "示例单位B")
            self.assertNotIn("归档单位(系统生成)", manual)
            overridden = map_piece(
                {"values": {"保管期限": "永久", "载体类型": "不存在的载体"}},
                meta["choices"],
                1,
            )
            self.assertEqual(overridden["保管期限"], "永久")
            self.assertEqual(overridden["载体类型"], "纸质")
            cleared, error = validate_defaults({"保管期限": ""}, meta["choices"])
            self.assertEqual(error, "")
            self.assertEqual(cleared["保管期限"], "")
            self.assertEqual(cleared["载体类型"], "纸质")
            omitted = map_piece(
                {"values": {}},
                meta["choices"],
                1,
                {"保管期限": "不存在", "载体类型": "照片"},
            )
            self.assertNotIn("保管期限", omitted)
            self.assertEqual(omitted["载体类型"], "照片")
            blanked = map_piece({"values": {}}, meta["choices"], 1, {"保管期限": ""})
            self.assertNotIn("保管期限", blanked)
            self.assertEqual(blanked["载体类型"], "纸质")
            rejected, error = validate_defaults(
                {"保管期限": "不存在", "载体类型": "纸质"}, meta["choices"]
            )
            self.assertIn("保管期限", error)
            self.assertIn("不存在", error)
            self.assertEqual(
                effective_defaults({"来源": "人工来源"})["来源"], "人工来源"
            )

    def test_default_sheet_prefers_piece_catalog(self):
        from openpyxl import Workbook

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "sheets.xlsx"
            book = Workbook()
            other = book.active
            other.title = "其他"
            other["A1"] = "甲"
            other["B1"] = "乙"
            piece = book.create_sheet("管理类_件")
            piece["A1"] = "题名"
            piece["B1"] = "责任者"
            book.save(path)
            book.close()
            self.assertEqual(inspect_template(path)["sheet"], "管理类_件")
            self.assertEqual(inspect_template(path, "其他")["sheet"], "其他")
            renamed = load_workbook(path)
            renamed["管理类_件"].title = "别的名字"
            renamed.save(path)
            renamed.close()
            with self.assertRaises(ValueError) as missing_sheet:
                inspect_template(path, "管理类_件")
            self.assertIn("管理类_件", str(missing_sheet.exception))
            self.assertIn("不会改用", str(missing_sheet.exception))
            self.assertEqual(format_file_size(5), "1")
            self.assertEqual(format_file_size(5, "MB"), "1")
            self.assertEqual(format_file_size(5, "KB"), "1")
            self.assertEqual(format_file_size(5, "B"), "5")
            self.assertEqual(format_file_size(20134807), "19")
            self.assertEqual(format_file_size(20134807, "MB"), "19")
            self.assertEqual(format_file_size(20134807, "KB"), "19663")
            self.assertLessEqual(int(format_file_size(20134807, "KB")), 9999999)
            self.assertGreater(int(format_file_size(20134807, "B")), 9999999)

    def test_summary_preserves_template_and_leaves_volumes_empty(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp = Path(tmp)
            first = tmp / "甲.png"
            second = tmp / "乙.png"
            first.write_bytes(b"ab")
            second.write_bytes(b"cdef")
            records = [
                {
                    "path": str(first),
                    "status": "待核对",
                    "export_order": 1,
                    "values": {
                        "题名": "第一份",
                        "责任者": "示例单位A",
                        "文件日期": "2026年10月26日",
                        "页数": "2",
                        "正本数量": "1",
                        "文号": "甲号",
                        "密级": "非涉密",
                    },
                },
                {
                    "path": str(second),
                    "status": "已核对",
                    "export_order": 4,
                    "values": {
                        "题名": "第二份",
                        "责任者": "没有这个单位",
                        "页数": "3",
                        "正本数量": "1",
                        "密级": "内部",
                    },
                },
            ]
            output = tmp / "汇总.xlsx"
            write_summary(output, records, TEMPLATE, "管理类_件")
            ok, lines = compare_workbooks(TEMPLATE, output)
            self.assertTrue(ok, "\n".join(lines))
            self.assertIn("结论：通过", lines)
            book = load_workbook(output)
            sheet = book["管理类_件"]
            self.assertEqual(sheet["D2"].value, "第一份")
            self.assertEqual(sheet["A2"].value, "甲.png")
            self.assertEqual(sheet["C2"].value, 1)
            self.assertEqual(sheet["F2"].value, "2026-10-26")
            self.assertEqual(sheet["G2"].value, "示例单位A")
            self.assertEqual(sheet["K2"].value, "非涉密")
            self.assertEqual(sheet["D3"].value, "第二份")
            self.assertEqual(sheet["C3"].value, 4)
            self.assertIsNone(sheet["K3"].value)
            self.assertIsNone(sheet["G3"].value)
            self.assertEqual(sheet["H2"].value, "30年")
            self.assertEqual(sheet["I2"].value, "纸质")
            self.assertEqual(sheet["H3"].value, "30年")
            self.assertIsNone(sheet["L2"].value)
            self.assertIsNone(sheet["D4"].value)
            self.assertTrue(all(cell.value is None for cell in book["管理类_卷"][2]))
            self.assertTrue(
                all(cell.value is None for cell in book["管理类_卷内件"][2])
            )
            book.close()
            separate = export_mapped_separate(records[:1], tmp, TEMPLATE, "管理类_件")
            self.assertFalse(separate[0]["error"])
            ok, lines = compare_workbooks(TEMPLATE, separate[0]["output"])
            self.assertTrue(ok, "\n".join(lines))

    def test_single_export_fills_a_header_only_or_self_closing_row(self):
        row_re = re.compile(r'<row\b[^>]*\br="2"(?!\d)[^>]*(?:/>|>.*?</row>)', re.S)

        def rewrite(path, transform):
            with zipfile.ZipFile(TEMPLATE) as source, zipfile.ZipFile(path, "w") as out:
                for info in source.infolist():
                    data = source.read(info.filename)
                    if info.filename == "xl/worksheets/sheet1.xml":
                        text, count = row_re.subn(
                            transform, data.decode("utf-8"), count=1
                        )
                        self.assertEqual(count, 1)
                        data = text.encode("utf-8")
                    out.writestr(info, data)

        with tempfile.TemporaryDirectory() as tmp:
            tmp = Path(tmp)
            scan = tmp / "丙.png"
            scan.write_bytes(b"xyz")
            record = {
                "path": str(scan),
                "status": "待核对",
                "export_order": 2,
                "values": {"题名": "缺行模板", "文件日期": "2026-10-26"},
            }
            for name, transform in (
                ("header-only.xlsx", lambda _match: ""),
                ("self-closing.xlsx", lambda _match: '<row r="2"/>'),
            ):
                template = tmp / name
                rewrite(template, transform)
                results = export_mapped_separate([record], tmp, template, "管理类_件")
                self.assertFalse(results[0]["error"], results[0]["error"])
                ok, lines = compare_workbooks(template, results[0]["output"])
                self.assertTrue(ok, "\n".join(lines))
                book = load_workbook(results[0]["output"])
                self.assertEqual(book["管理类_件"]["D2"].value, "缺行模板")
                self.assertEqual(book["管理类_件"]["C2"].value, 2)
                book.close()

    def test_compare_reports_a_structural_change(self):
        with tempfile.TemporaryDirectory() as tmp:
            broken = Path(tmp) / "broken.xlsx"
            with (
                zipfile.ZipFile(TEMPLATE) as source,
                zipfile.ZipFile(broken, "w") as output,
            ):
                for info in source.infolist():
                    data = source.read(info.filename)
                    if info.filename == "xl/workbook.xml":
                        data = data.replace("管理类_件".encode(), "管理类_改".encode())
                    output.writestr(info, data)
            ok, lines = compare_workbooks(TEMPLATE, broken)
            self.assertFalse(ok)
            self.assertTrue(any(line.startswith("未通过") for line in lines))
            self.assertIn("结论：未通过", lines[-1])

    def test_packaging_refuses_without_the_private_template(self):
        with tempfile.TemporaryDirectory() as tmp:
            with self.assertRaises(FileNotFoundError) as raised:
                require_standard_template(tmp)
            self.assertIn("无法打包", str(raised.exception))
            self.assertIn("standard_template.xlsx", str(raised.exception))
            present = Path(tmp) / "templates"
            present.mkdir()
            workbook = present / "standard_template.xlsx"
            workbook.write_bytes(b"placeholder")
            self.assertEqual(require_standard_template(tmp), workbook)
            missing = subprocess.run(
                [
                    sys.executable,
                    str(ROOT / "tools" / "require_template.py"),
                    tmp + "-absent",
                ],
                capture_output=True,
                text=True,
            )
            self.assertEqual(missing.returncode, 1)
            self.assertIn("无法打包", missing.stderr)
            found = subprocess.run(
                [sys.executable, str(ROOT / "tools" / "require_template.py"), tmp],
                capture_output=True,
                text=True,
            )
            self.assertEqual(found.returncode, 0)
            self.assertIn("已找到标准模板", found.stdout)


if __name__ == "__main__":
    unittest.main()
