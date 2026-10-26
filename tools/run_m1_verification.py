"""Headless M1 batch: import, OCR, export, compare, and retry failures.

Prints measured counts only. Exit status is 0 when the batch invariants hold.
"""

import csv
import io
import json
import os
import re
import sys
import tempfile
import time
from pathlib import Path

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tests"))

from openpyxl import load_workbook
from PySide6.QtCore import QTimer
from PySide6.QtWidgets import QApplication, QMessageBox, QPushButton

from app import Window
from fixtures import example_template
from template_io import inspect_template
from tools.compare_template import compare_workbooks
from tools.generate_scan_set import generate

EXTRACTED = [
    "题名",
    "责任者",
    "文件日期",
    "文件编号",
    "页数",
    "密级",
    "正本数量",
    "年度",
]
EXPORTED = EXTRACTED + [
    "归档单位",
    "归档单位(系统生成)",
    "保管期限",
    "载体类型",
    "电子文件数量",
    "电子文件关联项",
    "临时序号",
    "电子文件大小",
]
MANIFEST_COLUMNS = ("文件名", "状态", "原因", "对应Excel名")


def _cpu():
    text = Path("/proc/cpuinfo").read_text(encoding="utf-8", errors="replace")
    model = "unknown"
    count = 0
    for line in text.splitlines():
        if line.startswith("model name"):
            model = line.split(":", 1)[1].strip()
            count += 1
    if count == 0:
        count = os.cpu_count() or 0
    return model, count


def _cell(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _sheet_row(data):
    book = load_workbook(io.BytesIO(data))
    sheet = book["管理类_件"]
    headers = [_cell(cell.value) for cell in sheet[1]]
    values = {}
    for column, header in enumerate(headers, 1):
        if header:
            values[header] = _cell(sheet.cell(2, column).value)
    book.close()
    return values


def _compare_bytes(template, data):
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as handle:
        handle.write(data)
        path = handle.name
    try:
        return compare_workbooks(template, path)
    finally:
        Path(path).unlink(missing_ok=True)


class _Dialogs:
    def __init__(self):
        self.unreviewed = ""
        self.summaries = []
        self.buttons = []
        self.timer = QTimer()
        self.timer.timeout.connect(self._tick)
        self.timer.start(30)

    def _tick(self):
        box = QApplication.activeModalWidget()
        if not isinstance(box, QMessageBox) or getattr(box, "_handled", False):
            return
        box._handled = True
        title = box.windowTitle()
        if title == "尚有未核对文件":
            self.unreviewed = box.text()
            self._click(box, "继续导出")
        elif title == "批量导出结果":
            labels = [button.text() for button in box.findChildren(QPushButton)]
            self.buttons.append(labels)
            self.summaries.append(box.text())
            self._click(box, "关闭")
        else:
            self._click(box, "关闭")

    def _click(self, box, text):
        for button in box.findChildren(QPushButton):
            if button.text() == text:
                button.click()
                return
        box.reject()

    def stop(self):
        self.timer.stop()


def _wait_idle(app, window, timeout):
    deadline = time.monotonic() + timeout
    started = time.monotonic()
    while not window.busy() and time.monotonic() - started < 2:
        app.processEvents()
        time.sleep(0.01)
    while window.busy() and time.monotonic() < deadline:
        app.processEvents()
        time.sleep(0.02)
    app.processEvents()
    return not window.busy()


def _percent(hit, total):
    if total <= 0:
        return "0.00"
    return f"{100.0 * hit / total:.2f}"


def _format(report):
    lines = [
        f"输入 {report['n']}",
        f"可读 {report['readable']}",
        f"加密 {report['encrypted']}",
        f"损坏 {report['corrupt']}",
        f"倾斜 {report['skewed']}",
        f"模糊 {report['blurred']}",
        f"多页 {report['multipage']}",
        f"扫描外观 {report['scanned']}",
        f"可读页 {report['readable_pages']}",
        f"识别页 {report['recognized_pages']}",
        f"成功 {report['success']}",
        f"失败 {report['failure']}",
        f"成功+失败 {report['success_plus_failure']}",
        f"xlsx {report['xlsx']}",
        f"文件名一致 {report['name_match']}",
        f"compare通过 {report['compare_pass']}",
        f"compare未通过 {report['compare_fail']}",
        f"保管期限30年 {report['retention_ok']}",
        f"载体类型纸质 {report['media_ok']}",
        f"窗口默认保管期限30年 {report['window_retention']}",
        f"窗口默认载体类型纸质 {report['window_media']}",
        f"覆盖同名勾选 {report['overwrite_checked']}",
        f"清单行 {report['manifest_rows']}",
        f"清单列一致 {report['manifest_columns_ok']}",
        f"清单成功 {report['manifest_success']}",
        f"清单失败 {report['manifest_failure']}",
        f"结束对话框成功 {report['dialog_success']}",
        f"结束对话框失败 {report['dialog_failure']}",
        f"未核对对话框份数 {report['unreviewed_count']}",
        f"待核对记录 {report['review_count']}",
        f"加密原因命中 {report['encrypted_reason_hit']}",
        f"损坏原因命中 {report['corrupt_reason_hit']}",
        f"失败原因非空 {report['failure_reason_nonempty']}",
        f"重试按钮出现 {report['retry_button']}",
        f"重试文件 {report['retry_files']}",
        f"重试仅失败文件 {report['retry_only_failed']}",
        f"成功文件字节保持 {report['success_bytes_kept']}",
        f"重试后新增xlsx属于失败项 {report['retry_new_files_ok']}",
        f"识别秒 {report['ocr_seconds']:.2f}",
        f"导出秒 {report['export_seconds']:.2f}",
        f"重试秒 {report['retry_seconds']:.2f}",
        f"CPU {report['cpu_model']} × {report['cpu_count']}",
    ]
    for field in EXTRACTED:
        item = report["extracted"][field]
        lines.append(
            f"识别 {field} {item['hit']}/{item['total']} {_percent(item['hit'], item['total'])}"
        )
    for field in EXPORTED:
        item = report["exported"][field]
        lines.append(
            f"导出 {field} {item['hit']}/{item['total']} {_percent(item['hit'], item['total'])}"
        )
    for row in report["files"]:
        bits = " ".join(f"{name}{value}" for name, value in row["extracted"].items())
        lines.append(
            (
                f"文件 {row['filename']} 状态 {row['status']} 页 {row['pages']} "
                f"xlsx {row['excel']} 原因命中 {row['reason_hit']} {bits}"
            ).rstrip()
        )
    lines.append("遗留")
    if report["issues"]:
        lines.extend(report["issues"])
    else:
        lines.append("无")
    return "\n".join(lines)


def run():
    app = QApplication.instance() or QApplication([])
    problems = []
    cpu_model, cpu_count = _cpu()
    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        records = generate(tmp / "in")
        output = tmp / "out"
        output.mkdir()
        remembered = tmp / "remembered"
        remembered.mkdir()
        template = example_template()
        meta = inspect_template(template, "管理类_件")
        window = Window(tmp / "project.json")
        window.use_template = True
        window.template_notice = None
        window.template_path = str(template)
        window.template_sheet = "管理类_件"
        window.fields = meta["headers"]
        window.choices = meta["choices"]
        window.build_form()
        window._apply_export_enabled()
        window.export_folder = str(remembered)
        paths = [str(tmp / "in" / record["filename"]) for record in records]
        window.add_paths(paths)
        export_calls = []
        original_start = window._start_export

        def track_start(indexes):
            export_calls.append(list(indexes))
            return original_start(indexes)

        window._start_export = track_start
        ocr_started = time.monotonic()
        window.start()
        if not _wait_idle(app, window, 300):
            problems.append("识别超时")
        ocr_seconds = time.monotonic() - ocr_started
        snapshot = []
        for record in window.records:
            snapshot.append(
                {
                    "status": record.get("status", ""),
                    "warning": record.get("warning", ""),
                    "pages": record.get("pages", 0),
                    "values": dict(record.get("values") or {}),
                }
            )
        dialogs = _Dialogs()
        export_started = time.monotonic()
        from unittest.mock import patch

        with patch(
            "app.QFileDialog.getExistingDirectory", return_value=str(output)
        ) as choose:
            window.export_btn.click()
            deadline = time.monotonic() + 120
            while len(dialogs.summaries) < 1 and time.monotonic() < deadline:
                app.processEvents()
                time.sleep(0.02)
            _wait_idle(app, window, 60)
            export_seconds = time.monotonic() - export_started
            manifest_path = output / "处理清单.csv"
            first_manifest = (
                manifest_path.read_bytes() if manifest_path.is_file() else b""
            )
            before = {path.name: path.read_bytes() for path in output.glob("*.xlsx")}
            retry_started = time.monotonic()
            window.retry_failed()
            deadline = time.monotonic() + 300
            while len(dialogs.summaries) < 2 and time.monotonic() < deadline:
                app.processEvents()
                time.sleep(0.02)
            if not _wait_idle(app, window, 60):
                problems.append("重试超时")
            retry_seconds = time.monotonic() - retry_started
        folder_arg = choose.call_args_list[0].args[2] if choose.call_args_list else ""
        if folder_arg != str(remembered):
            problems.append("未从上次文件夹开始选择")
        dialogs.stop()
        after = {path.name: path.read_bytes() for path in output.glob("*.xlsx")}
        second_manifest = manifest_path.read_bytes() if manifest_path.is_file() else b""
        window.dirty = False
        window.close()

    return _assemble(
        records,
        snapshot,
        problems,
        dialogs,
        export_calls,
        first_manifest,
        second_manifest,
        before,
        after,
        template,
        ocr_seconds,
        export_seconds,
        retry_seconds,
        cpu_model,
        cpu_count,
        window_defaults={
            "保管期限": window.catalog_defaults.get("保管期限", ""),
            "载体类型": window.catalog_defaults.get("载体类型", ""),
            "overwrite": int(window.overwrite_same_name),
        },
    )


def _assemble(
    records,
    snapshot,
    problems,
    dialogs,
    export_calls,
    first_manifest,
    second_manifest,
    before,
    after,
    template,
    ocr_seconds,
    export_seconds,
    retry_seconds,
    cpu_model,
    cpu_count,
    window_defaults,
):
    issues = []
    columns, rows = (), []
    if not first_manifest:
        problems.append("没有处理清单")
    else:
        reader = csv.DictReader(io.StringIO(first_manifest.decode("utf-8-sig")))
        columns = tuple(reader.fieldnames or ())
        rows = list(reader)
    if not dialogs.summaries:
        problems.append("没有结束对话框")
    success_rows = [row for row in rows if row.get("状态") == "成功"]
    failure_rows = [row for row in rows if row.get("状态") == "失败"]
    if columns != MANIFEST_COLUMNS:
        problems.append("处理清单列不一致")
    if len(rows) != len(records):
        problems.append("处理清单行数不等于输入")
    if len(success_rows) + len(failure_rows) != len(records):
        problems.append("成功加失败不等于输入")
    if len(before) != len(success_rows):
        problems.append("xlsx数量不等于成功份数")

    name_match = 0
    compare_pass = 0
    compare_fail = 0
    retention_ok = 0
    media_ok = 0
    files = []
    extracted_hits = {field: 0 for field in EXTRACTED}
    exported_hits = {field: 0 for field in EXPORTED}
    readable = [record for record in records if not record["expect_failure"]]
    recognized_pages = 0
    for index, record in enumerate(records):
        seen = snapshot[index] if index < len(snapshot) else {}
        values = seen.get("values") or {}
        manifest = rows[index] if index < len(rows) else {}
        reason = manifest.get("原因", "")
        status = manifest.get("状态", "")
        excel = manifest.get("对应Excel名", "")
        reason_hit = 0
        if record["expect_failure"]:
            reason_hit = int(record["expect_failure"] in reason and status == "失败")
            if not reason_hit:
                problems.append(record["filename"] + "失败原因不符")
        elif status == "成功":
            expected_name = Path(record["filename"]).with_suffix(".xlsx").name
            if excel == expected_name and excel in before:
                name_match += 1
            else:
                problems.append(record["filename"] + "文件名不一致")
            sheet = _sheet_row(before[excel]) if excel in before else {}
            if sheet.get("保管期限") == "30年":
                retention_ok += 1
            else:
                problems.append(record["filename"] + "保管期限不是30年")
            if sheet.get("载体类型") == "纸质":
                media_ok += 1
            else:
                problems.append(record["filename"] + "载体类型不是纸质")
            ok, compare_lines = (
                _compare_bytes(template, before[excel])
                if excel in before
                else (False, [])
            )
            if ok:
                compare_pass += 1
            else:
                compare_fail += 1
                problems.append(record["filename"] + "格式比较未通过")
            for field in EXPORTED:
                if sheet.get(field, "") == record["fields"].get(field, ""):
                    exported_hits[field] += 1
        field_bits = {}
        if not record["expect_failure"]:
            if seen.get("pages"):
                recognized_pages += int(seen.get("pages") or 0)
            for field in EXTRACTED:
                hit = int(
                    str(values.get(field, "")).strip()
                    == record["fields"].get(field, "")
                )
                field_bits[field] = hit
                extracted_hits[field] += hit
            if status != "成功":
                issues.append(record["filename"] + "识别或导出失败")
        files.append(
            {
                "filename": record["filename"],
                "status": status or seen.get("status", ""),
                "pages": seen.get("pages", record["pages"]),
                "excel": excel,
                "reason_hit": reason_hit,
                "extracted": field_bits,
            }
        )

    failure_reason_nonempty = sum(
        1 for row in failure_rows if row.get("原因", "").strip()
    )
    if failure_reason_nonempty != len(failure_rows):
        problems.append("有失败行没有原因")
    dialog_success, dialog_failure = _counts(
        dialogs.summaries[0] if dialogs.summaries else ""
    )
    if dialogs.summaries and (dialog_success, dialog_failure) != (
        len(success_rows),
        len(failure_rows),
    ):
        problems.append("结束对话框计数与清单不一致")
    unreviewed = _unreviewed(dialogs.unreviewed)
    review_count = sum(item.get("status") == "待核对" for item in snapshot)
    if review_count and unreviewed != review_count:
        problems.append("未核对对话框份数不一致")
    retry_labels = dialogs.buttons[0] if dialogs.buttons else []
    retry_button = int("重试失败文件" in retry_labels)
    if failure_rows and not retry_button:
        problems.append("结束对话框没有重试失败文件")
    failed_indexes = [
        index for index, row in enumerate(rows) if row.get("状态") == "失败"
    ]
    retry_indexes = export_calls[1] if len(export_calls) > 1 else []
    retry_only_failed = int(retry_indexes == failed_indexes and len(export_calls) == 2)
    if not retry_only_failed:
        problems.append("重试没有只处理失败文件")
    success_bytes_kept = sum(
        1 for name, data in before.items() if after.get(name) == data
    )
    if success_bytes_kept != len(before):
        problems.append("重试改写了已成功的文件")
    failed_names = {row["文件名"] for row in failure_rows}
    new_names = set(after) - set(before)
    retry_new_ok = int(
        all(
            Path(name).name
            in {Path(item).with_suffix(".xlsx").name for item in failed_names}
            for name in new_names
        )
    )
    if not retry_new_ok:
        problems.append("重试写出了失败项以外的文件")
    if second_manifest:
        second_text = io.StringIO(second_manifest.decode("utf-8-sig"))
        second_rows = list(csv.DictReader(second_text))
        second_names = {row["文件名"] for row in second_rows}
        if second_names != failed_names or len(second_rows) != len(failed_names):
            problems.append("重试后的处理清单不是失败文件")
    else:
        problems.append("重试后没有处理清单")

    for field, hit in extracted_hits.items():
        if hit != len(readable):
            missed = [
                item["filename"]
                for item in files
                if item["extracted"].get(field, 1) == 0
            ]
            issues.append(f"{field}不一致 {len(readable) - hit} " + " ".join(missed))
    if window_defaults["保管期限"] != "30年" or window_defaults["载体类型"] != "纸质":
        problems.append("窗口默认值不是30年和纸质")
    issues.append("并行OCR评估未进入本分支")

    report = {
        "ok": not problems,
        "problems": problems,
        "issues": issues,
        "n": len(records),
        "readable": len(readable),
        "encrypted": sum("加密" in record["traits"] for record in records),
        "corrupt": sum("损坏" in record["traits"] for record in records),
        "skewed": sum("倾斜" in record["traits"] for record in records),
        "blurred": sum("模糊" in record["traits"] for record in records),
        "multipage": sum("多页" in record["traits"] for record in records),
        "scanned": sum("扫描" in record["traits"] for record in records),
        "readable_pages": sum(record["pages"] for record in readable),
        "recognized_pages": recognized_pages,
        "success": len(success_rows),
        "failure": len(failure_rows),
        "success_plus_failure": len(success_rows) + len(failure_rows),
        "xlsx": len(before),
        "name_match": name_match,
        "compare_pass": compare_pass,
        "compare_fail": compare_fail,
        "retention_ok": retention_ok,
        "media_ok": media_ok,
        "window_retention": int(window_defaults["保管期限"] == "30年"),
        "window_media": int(window_defaults["载体类型"] == "纸质"),
        "overwrite_checked": window_defaults["overwrite"],
        "manifest_rows": len(rows),
        "manifest_columns_ok": int(columns == MANIFEST_COLUMNS),
        "manifest_success": len(success_rows),
        "manifest_failure": len(failure_rows),
        "dialog_success": dialog_success,
        "dialog_failure": dialog_failure,
        "unreviewed_count": unreviewed,
        "review_count": review_count,
        "encrypted_reason_hit": sum(
            item["reason_hit"] for item in files if "加密" in item["filename"]
        ),
        "corrupt_reason_hit": sum(
            item["reason_hit"] for item in files if "损坏" in item["filename"]
        ),
        "failure_reason_nonempty": failure_reason_nonempty,
        "retry_button": retry_button,
        "retry_files": len(retry_indexes),
        "retry_only_failed": retry_only_failed,
        "success_bytes_kept": success_bytes_kept,
        "retry_new_files_ok": retry_new_ok,
        "ocr_seconds": ocr_seconds,
        "export_seconds": export_seconds,
        "retry_seconds": retry_seconds,
        "cpu_model": cpu_model,
        "cpu_count": cpu_count,
        "extracted": {
            field: {"hit": extracted_hits[field], "total": len(readable)}
            for field in EXTRACTED
        },
        "exported": {
            field: {"hit": exported_hits[field], "total": len(success_rows)}
            for field in EXPORTED
        },
        "files": files,
    }
    report["text"] = _format(report)
    return report


def _counts(summary):
    match = re.search(r"成功\s*(\d+)\s*份，失败\s*(\d+)\s*份", summary)
    if not match:
        return -1, -1
    return int(match.group(1)), int(match.group(2))


def _unreviewed(text):
    match = re.search(r"有\s*(\d+)\s*份尚未核对", text or "")
    if not match:
        return -1
    return int(match.group(1))


if __name__ == "__main__":
    result = run()
    print(result["text"])
    if result["problems"]:
        print("问题")
        print("\n".join(result["problems"]))
    sys.exit(0 if result["ok"] else 1)
