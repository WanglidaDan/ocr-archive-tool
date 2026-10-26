"""Windows application XLSX template adapter, preserving the original schema."""

from datetime import datetime
import math
from pathlib import Path
from copy import copy
from catalog_defaults import DEFAULT_SHEET
from extraction import canonical, normalize_date

NUMERIC = {
    "临时序号",
    "序号",
    "页数",
    "正本数量",
    "数量",
    "年度",
    "件数",
    "电子文件大小",
    "电子文件数量",
}
DATES = {"文件日期", "日期", "成文日期", "形成日期", "发文日期", "起日期", "止日期"}


def _numeric_bound(formula):
    """Return a plain numeric limit, or None when the bound is not a number.

    Data validations sometimes store formula1/formula2 as a formula or a cell
    reference (TODAY(), $A$1). Those cannot be compared with float(); skip the
    bound instead of failing the export.
    """
    if not formula:
        return None
    try:
        return float(formula)
    except (TypeError, ValueError):
        return None


def _outside_bounds(value, formula1, formula2):
    low = _numeric_bound(formula1)
    high = _numeric_bound(formula2)
    return (low is not None and value < low) or (high is not None and value > high)


def standard_value(field, value, choices):
    # Preserve the provided dictionary: do not invent a new classification item.
    if (
        field == "密级"
        and value == "内部"
        and value not in choices
        and "内部资料" in choices
    ):
        return "内部资料"
    return value


def inspect_template(path, sheet=None):
    from openpyxl import load_workbook

    book = load_workbook(path)
    names = [s.title for s in book if s.max_column > 1 and s.cell(1, 1).value]
    if sheet in names:
        selected = sheet
    elif DEFAULT_SHEET in names:
        selected = DEFAULT_SHEET
    else:
        selected = names[0] if names else None
    if selected is None:
        raise ValueError("没有找到包含表头的目录工作表")
    ws = book[selected]
    headers = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
    if any(not h for h in headers) or len(set(headers)) != len(headers):
        raise ValueError("目录表头不能为空或重复")
    choices = {}
    for dv in ws.data_validations.dataValidation:
        if dv.type != "list":
            continue
        cols = [c.column for c in ws[2] if c.coordinate in dv.sqref]
        if not cols:
            continue
        formula = (dv.formula1 or "").lstrip("=")
        options = []
        if formula.startswith('"') and formula.endswith('"'):
            options = formula[1:-1].split(",")
        elif formula in book.defined_names:
            for title, cell_range in book.defined_names[formula].destinations:
                options.extend(
                    str(c.value)
                    for row in book[title][cell_range]
                    for c in row
                    if c.value is not None
                )
        for col in cols:
            if options:
                choices[headers[col - 1]] = options
    result = {
        "sheets": names,
        "sheet": selected,
        "headers": headers,
        "choices": choices,
        "all_sheets": book.sheetnames,
    }
    book.close()
    return result


def _patches_for_row(meta, book, ws, fields, row):
    """Turn one record into cell patches. Dates stay ISO text to keep styles.xml."""
    import re
    from openpyxl.utils import get_column_letter
    from openpyxl.utils.datetime import to_excel

    patches = {}
    for col, field in enumerate(meta["headers"], 1):
        value = fields.get(field, "")
        if value is None or not str(value).strip():
            continue
        value = standard_value(
            field, str(value).strip(), meta["choices"].get(field, [])
        )
        address = f"{get_column_letter(col)}{row}"
        if field in meta["choices"] and value not in meta["choices"][field]:
            raise ValueError(field + "不在原模板字典内，请核对；未改动字典")
        if field in DATES:
            value, error = normalize_date(value)
            if error:
                raise ValueError(field + "：" + error)
            serial = to_excel(datetime.fromisoformat(value), book.epoch)
            for dv in ws.data_validations.dataValidation:
                if dv.type == "date" and address in dv.sqref:
                    if _outside_bounds(serial, dv.formula1, dv.formula2):
                        raise ValueError(field + "超出原模板日期范围")
            patches[address] = ("inlineStr", value)
        elif field in NUMERIC:
            try:
                number = float(value)
            except ValueError:
                raise ValueError(field + "必须是数字")
            if (
                not math.isfinite(number)
                or number < 0
                or (field != "电子文件大小" and not number.is_integer())
            ):
                raise ValueError(field + "必须是有效非负数值")
            for dv in ws.data_validations.dataValidation:
                if dv.type in ("whole", "decimal") and address in dv.sqref:
                    if _outside_bounds(number, dv.formula1, dv.formula2):
                        raise ValueError(field + "超出原模板范围")
            patches[address] = (
                "n",
                str(int(number)) if number.is_integer() else str(number),
            )
        else:
            patches[address] = ("inlineStr", value)
        kind_value = patches[address][1]
        if len(kind_value) > 32767 or re.search(
            r"[\x00-\x08\x0b\x0c\x0e-\x1f\ufffe\uffff]", kind_value
        ):
            raise ValueError("字段含非法字符或超过Excel单元格长度")
    return patches


def _apply_row_patches(rowxml, patches):
    import re
    from xml.sax.saxutils import escape
    from openpyxl.utils import column_index_from_string

    for address, (kind, value) in patches.items():
        body = (
            ('<is><t xml:space="preserve">' + escape(value) + "</t></is>")
            if kind == "inlineStr"
            else "<v>" + value + "</v>"
        )
        cellpattern = re.compile(
            r'<c\b(?P<attrs>[^>]*\br="'
            + address
            + r'"[^>]*?)(?:\s*/>|>(?P<body>.*?)</c>)',
            re.S,
        )
        found = cellpattern.search(rowxml)
        if found:
            attrs = re.sub(r'\s+t="[^"]*"', "", found["attrs"]).rstrip()
            cell = "<c" + attrs + ' t="' + kind + '">' + body + "</c>"
            rowxml = rowxml[: found.start()] + cell + rowxml[found.end() :]
        else:
            cell = '<c r="' + address + '" t="' + kind + '">' + body + "</c>"
            colnum = column_index_from_string(re.match(r"[A-Z]+", address)[0])
            later = next(
                (
                    item
                    for item in re.finditer(r'<c\b[^>]*\br="([A-Z]+)\d+"', rowxml)
                    if column_index_from_string(item[1]) > colnum
                ),
                None,
            )
            offset = later.start() if later else rowxml.rfind("</row>")
            rowxml = rowxml[:offset] + cell + rowxml[offset:]
    return rowxml


def _write_patches(template_path, sheet_name, patches_by_row, output_path):
    """Copy the template and replace only the listed data cells."""
    import re
    import posixpath
    import zipfile
    import xml.etree.ElementTree as ET

    ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    ridns = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
    with zipfile.ZipFile(template_path) as source:
        workbook = ET.fromstring(source.read("xl/workbook.xml"))
        node = next(
            item
            for item in workbook.find("s:sheets", ns)
            if item.attrib["name"] == sheet_name
        )
        rels = ET.fromstring(source.read("xl/_rels/workbook.xml.rels"))
        target = next(
            item.attrib["Target"]
            for item in rels
            if item.attrib["Id"] == node.attrib[ridns]
        )
        member = (
            target.lstrip("/")
            if target.startswith("/")
            else posixpath.normpath("xl/" + target)
        )
        xml = source.read(member).decode("utf-8")
        for row in sorted(patches_by_row):
            patches = patches_by_row[row]
            if not patches:
                continue
            rowpattern = re.compile(
                r'<row\b[^>]*\br="' + str(row) + r'"(?!\d)[^>]*(?:/>|>.*?</row>)',
                re.S,
            )
            rowmatch = rowpattern.search(xml)
            if rowmatch and rowmatch.group(0).endswith("/>"):
                expanded = rowmatch.group(0)[:-2].rstrip() + "></row>"
                xml = xml[: rowmatch.start()] + expanded + xml[rowmatch.end() :]
                rowmatch = rowpattern.search(xml)
            if not rowmatch:
                inserted = '<row r="' + str(row) + '"></row>'
                later = next(
                    (
                        item
                        for item in re.finditer(r'<row\b[^>]*\br="(\d+)"', xml)
                        if int(item.group(1)) > row
                    ),
                    None,
                )
                if later is not None:
                    xml = xml[: later.start()] + inserted + xml[later.start() :]
                else:
                    end = xml.rfind("</sheetData>")
                    if end < 0:
                        raise ValueError("模板数据行结构不受支持；未重新构建工作表")
                    xml = xml[:end] + inserted + xml[end:]
                rowmatch = rowpattern.search(xml)
                if not rowmatch:
                    raise ValueError("模板数据行结构不受支持；未重新构建工作表")
            rowxml = _apply_row_patches(rowmatch[0], patches)
            xml = xml[: rowmatch.start()] + rowxml + xml[rowmatch.end() :]
        with zipfile.ZipFile(output_path, "w") as output:
            for info in source.infolist():
                output.writestr(
                    info,
                    xml.encode("utf-8")
                    if info.filename == member
                    else source.read(info.filename),
                )


def _load_sheet(template_path, sheet):
    from openpyxl import load_workbook

    if not Path(template_path).is_file():
        raise ValueError("找不到模板：" + str(template_path))
    meta = inspect_template(template_path, sheet)
    book = load_workbook(template_path)
    return meta, book, book[meta["sheet"]]


def export_template(path, records, template_path, sheet):
    """Patch one data row. Copy every other OOXML member byte-for-byte."""
    if len(records) != 1:
        raise ValueError("原模板模式每个Excel仅对应一份文件，请使用批量分开导出")
    if Path(path).resolve() == Path(template_path).resolve():
        raise ValueError("不能覆盖原模板")
    meta, book, ws = _load_sheet(template_path, sheet)
    row = 2
    while row <= ws.max_row and any(
        ws.cell(row, column).value is not None
        for column in range(1, len(meta["headers"]) + 1)
    ):
        row += 1
    patches = _patches_for_row(meta, book, ws, records[0].get("values", {}), row)
    book.close()
    _write_patches(template_path, meta["sheet"], {row: patches}, path)


def export_rows(path, records, template_path, sheet, start_row=2):
    """Write records into consecutive rows of one template copy, starting at row 2."""
    if not records:
        raise ValueError("没有可写入的记录")
    if Path(path).resolve() == Path(template_path).resolve():
        raise ValueError("不能覆盖原模板")
    meta, book, ws = _load_sheet(template_path, sheet)
    patches_by_row = {}
    try:
        for offset, record in enumerate(records):
            row = start_row + offset
            patches_by_row[row] = _patches_for_row(
                meta, book, ws, record.get("values", {}), row
            )
    finally:
        book.close()
    _write_patches(template_path, meta["sheet"], patches_by_row, path)
